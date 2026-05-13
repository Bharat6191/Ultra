"""Global executive / operational intelligence for the main workspace dashboard."""

from __future__ import annotations

from calendar import monthrange
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from modules.approvals.model import ApprovalTask
from modules.contractor.models import Contractor, ContractorDocument
from modules.contractor_rates.models import ContractorRate
from modules.invoices.models import Invoice
from modules.org_units.model import OrgUnit
from modules.part_master.models import PartMaster
from modules.work_orders.models import WorkOrder, WorkOrderItem


def _d0(v: Decimal | None) -> Decimal:
    return v if v is not None else Decimal("0")


def _month_range(year: int, month: int) -> tuple[datetime, datetime]:
    start = datetime(year, month, 1, 0, 0, 0, tzinfo=timezone.utc)
    last = monthrange(year, month)[1]
    end = datetime(year, month, last, 23, 59, 59, 999999, tzinfo=timezone.utc) + timedelta(days=1)
    return start, end


def _iter_last_n_months(n: int) -> list[tuple[int, int]]:
    out: list[tuple[int, int]] = []
    d = date.today().replace(day=1)
    for _ in range(n):
        out.append((d.year, d.month))
        if d.month == 1:
            d = d.replace(year=d.year - 1, month=12)
        else:
            d = d.replace(month=d.month - 1)
    out.reverse()
    return out


@dataclass
class IntelligenceFilters:
    date_from: date | None = None
    date_to: date | None = None
    plant_id: int | None = None
    contractor_id: int | None = None
    work_order_status: str | None = None
    negotiation_status: str | None = None
    invoice_status: str | None = None
    part_pricing_method: str | None = None


class DashboardIntelligenceService:
    def __init__(self, db: Session) -> None:
        self._db = db

    def _wo_base(self, flt: IntelligenceFilters):
        q = select(WorkOrder)
        if flt.plant_id:
            q = q.where(WorkOrder.org_unit_id == int(flt.plant_id))
        if flt.contractor_id:
            q = q.where(WorkOrder.contractor_id == int(flt.contractor_id))
        if flt.date_from:
            q = q.where(WorkOrder.created_at >= datetime.combine(flt.date_from, datetime.min.time()).replace(tzinfo=timezone.utc))
        if flt.date_to:
            end = datetime.combine(flt.date_to + timedelta(days=1), datetime.min.time()).replace(tzinfo=timezone.utc)
            q = q.where(WorkOrder.created_at < end)
        if flt.work_order_status:
            q = q.where(WorkOrder.status == str(flt.work_order_status).strip())
        return q

    def _inv_base(self, flt: IntelligenceFilters):
        q = select(Invoice)
        if flt.plant_id:
            q = q.where(Invoice.org_unit_id == int(flt.plant_id))
        if flt.contractor_id:
            q = q.where(Invoice.contractor_id == int(flt.contractor_id))
        if flt.date_from:
            q = q.where(Invoice.invoice_date >= flt.date_from)
        if flt.date_to:
            q = q.where(Invoice.invoice_date <= flt.date_to)
        if flt.invoice_status:
            q = q.where(Invoice.status == str(flt.invoice_status).strip())
        return q

    def _rates_base(self, flt: IntelligenceFilters):
        q = select(ContractorRate)
        if flt.contractor_id:
            q = q.where(ContractorRate.contractor_id == int(flt.contractor_id))
        if flt.negotiation_status:
            q = q.where(ContractorRate.status == str(flt.negotiation_status).strip())
        return q

    def build(self, grants: set[str], flt: IntelligenceFilters) -> dict[str, Any]:
        def has(*codes: str) -> bool:
            for c in codes:
                variants = {c, c.replace(".", ":", 1) if "." in c else c, c.replace(":", ".", 1) if ":" in c else c}
                if not grants.isdisjoint(variants):
                    return True
            return False

        out: dict[str, Any] = {"enabled": True, "filters": {k: (v.isoformat() if isinstance(v, date) else v) for k, v in asdict(flt).items()}}

        if not has(
            "contractor.view",
            "work_orders.view",
            "work_orders.create",
            "invoices.view",
            "invoices.create",
            "contractor_rates.view",
            "approval.view",
            "task.view",
            "users.view",
        ):
            return {"enabled": False, "message": "No dashboard intelligence modules for your role."}

        # --- Executive ---
        exec_block: dict[str, Any] = {}
        if has("contractor.view"):
            exec_block["active_contractors"] = int(
                self._db.scalar(select(func.count()).select_from(Contractor).where(Contractor.status == "active")) or 0
            )
            today = date.today()
            warn_until = today + timedelta(days=30)
            exec_block["expiring_rates_30d"] = int(
                self._db.scalar(
                    select(func.count())
                    .select_from(ContractorRate)
                    .where(
                        ContractorRate.status == "approved",
                        ContractorRate.effective_to.isnot(None),
                        ContractorRate.effective_to >= today,
                        ContractorRate.effective_to <= warn_until,
                    )
                )
                or 0
            )
        if has("work_orders.view", "work_orders.create"):
            wq = self._wo_base(flt)
            rows = list(self._db.scalars(wq).all())
            by_st = {}
            for w in rows:
                by_st[str(w.status).lower()] = by_st.get(str(w.status).lower(), 0) + 1
            exec_block["work_orders_total"] = len(rows)
            exec_block["work_orders_active"] = by_st.get("active", 0) + by_st.get("approved", 0)
            exec_block["work_orders_pending"] = by_st.get("draft", 0) + by_st.get("pending_approval", 0)
            exec_block["work_orders_completed"] = by_st.get("closed", 0)
        if has("invoices.view", "invoices.create"):
            iq = self._inv_base(flt)
            invs = list(self._db.scalars(iq).all())
            exec_block["invoice_total_value"] = float(sum(_d0(i.total_amount) for i in invs))
            pend_st = ("draft", "submitted", "blocked", "pending_exception_approval")
            exec_block["invoice_pending_value"] = float(
                sum(_d0(i.total_amount) for i in invs if str(i.status).lower() in pend_st)
            )
            exec_block["invoice_approved_value"] = float(
                sum(_d0(i.total_amount) for i in invs if str(i.status).lower() in ("approved", "paid"))
            )
            exec_block["invoice_rejected_value"] = float(
                sum(_d0(i.total_amount) for i in invs if str(i.status).lower() == "rejected")
            )
        if has("contractor_rates.view"):
            rq = self._rates_base(flt)
            rates = list(self._db.scalars(rq).all())
            by_r = {}
            for r in rates:
                by_r[str(r.status).lower()] = by_r.get(str(r.status).lower(), 0) + 1
            exec_block["negotiations_total"] = len(rates)
            exec_block["negotiations_pending"] = by_r.get("pending_approval", 0) + by_r.get("draft", 0)
            sav_q = select(func.coalesce(func.sum(ContractorRate.savings_amount), 0)).where(
                ContractorRate.status == "approved",
                ContractorRate.savings_amount.isnot(None),
            )
            if flt.contractor_id:
                sav_q = sav_q.where(ContractorRate.contractor_id == int(flt.contractor_id))
            sav = self._db.scalar(sav_q) or 0
            exec_block["negotiation_savings_total"] = float(sav)
            prem = Decimal("0")
            try:
                pq = (
                    select(ContractorRate.negotiated_rate, PartMaster.base_rate)
                    .join(PartMaster, PartMaster.id == ContractorRate.part_master_id)
                    .where(ContractorRate.status == "approved")
                )
                if flt.contractor_id:
                    pq = pq.where(ContractorRate.contractor_id == int(flt.contractor_id))
                for neg, base in self._db.execute(pq).all():
                    if neg is None or base is None:
                        continue
                    d = Decimal(str(neg)) - Decimal(str(base))
                    if d > 0:
                        prem += d
            except Exception:
                pass
            exec_block["premium_above_base_total"] = float(prem)

        # Pending approvals (tasks inbox — same visibility as summary is complex; use global pending count for approvers)
        if has("approval.view", "task.view"):
            exec_block["pending_approval_tasks"] = int(
                self._db.scalar(
                    select(func.count()).select_from(ApprovalTask).where(ApprovalTask.status == "pending")
                )
                or 0
            )

        # Month / prev month spend & savings (invoice date month; savings by approved_at month)
        cy, cm = date.today().year, date.today().month
        if cm == 1:
            py, pm = cy - 1, 12
        else:
            py, pm = cy, cm - 1
        cur_s, cur_e = _month_range(cy, cm)
        prev_s, prev_e = _month_range(py, pm)

        def inv_sum_range(start: datetime, end: datetime) -> float:
            q = select(func.coalesce(func.sum(Invoice.total_amount), 0)).where(
                Invoice.invoice_date >= start.date(),
                Invoice.invoice_date < end.date(),
                Invoice.status.in_(("approved", "paid")),
            )
            if flt.plant_id:
                q = q.where(Invoice.org_unit_id == int(flt.plant_id))
            if flt.contractor_id:
                q = q.where(Invoice.contractor_id == int(flt.contractor_id))
            return float(self._db.scalar(q) or 0)

        exec_block["monthly_spend_current"] = inv_sum_range(cur_s, cur_e)
        exec_block["monthly_spend_previous"] = inv_sum_range(prev_s, prev_e)

        def savings_month(start: datetime, end: datetime) -> float:
            q = select(func.coalesce(func.sum(ContractorRate.savings_amount), 0)).where(
                ContractorRate.status == "approved",
                ContractorRate.savings_amount.isnot(None),
                ContractorRate.approved_at.isnot(None),
                ContractorRate.approved_at >= start,
                ContractorRate.approved_at < end,
            )
            if flt.contractor_id:
                q = q.where(ContractorRate.contractor_id == int(flt.contractor_id))
            return float(self._db.scalar(q) or 0)

        exec_block["monthly_savings_current"] = savings_month(cur_s, cur_e)
        exec_block["monthly_savings_previous"] = savings_month(prev_s, prev_e)

        # Sparklines last 6 months — WO created, invoice spend, negotiations created
        spark_wo: list[dict[str, Any]] = []
        spark_inv: list[dict[str, Any]] = []
        spark_neg: list[dict[str, Any]] = []
        for y, m in _iter_last_n_months(6):
            s, e = _month_range(y, m)
            key = f"{y}-{m:02d}"
            wct = int(
                self._db.scalar(
                    select(func.count())
                    .select_from(WorkOrder)
                    .where(WorkOrder.created_at >= s, WorkOrder.created_at < e)
                )
                or 0
            )
            spark_wo.append({"period": key, "value": wct})
            inv_amt = float(
                self._db.scalar(
                    select(func.coalesce(func.sum(Invoice.total_amount), 0)).where(
                        Invoice.invoice_date >= s.date(),
                        Invoice.invoice_date < e.date(),
                        Invoice.status.in_(("approved", "paid")),
                    )
                )
                or 0
            )
            spark_inv.append({"period": key, "value": inv_amt})
            nct = int(
                self._db.scalar(
                    select(func.count())
                    .select_from(ContractorRate)
                    .where(ContractorRate.created_at >= s, ContractorRate.created_at < e)
                )
                or 0
            )
            spark_neg.append({"period": key, "value": nct})

        exec_block["sparklines"] = {
            "work_orders_created": spark_wo,
            "invoice_spend": spark_inv,
            "negotiations_created": spark_neg,
        }
        out["executive"] = exec_block

        # --- Operations (work orders) ---
        if has("work_orders.view", "work_orders.create"):
            wrows = list(self._db.scalars(self._wo_base(flt)).all())
            wo_sub = self._wo_base(flt).subquery()
            by_st = {}
            for w in wrows:
                by_st[str(w.status).lower()] = by_st.get(str(w.status).lower(), 0) + 1
            donut = [{"name": k, "value": v} for k, v in sorted(by_st.items(), key=lambda x: -x[1])]
            # monthly created / completed (last 8 months, unfiltered by flt for trend shape)
            months = _iter_last_n_months(8)
            wo_trend = []
            for y, m in months:
                s, e = _month_range(y, m)
                c_created = int(
                    self._db.scalar(
                        select(func.count()).select_from(WorkOrder).where(WorkOrder.created_at >= s, WorkOrder.created_at < e)
                    )
                    or 0
                )
                c_closed = int(
                    self._db.scalar(
                        select(func.count())
                        .select_from(WorkOrder)
                        .where(WorkOrder.status == "closed", WorkOrder.updated_at >= s, WorkOrder.updated_at < e)
                    )
                    or 0
                )
                wo_trend.append({"month": f"{y}-{m:02d}", "created": c_created, "completed": c_closed})
            plant_rows = self._db.execute(
                select(OrgUnit.name, func.count(wo_sub.c.id))
                .select_from(wo_sub)
                .join(OrgUnit, OrgUnit.id == wo_sub.c.org_unit_id)
                .group_by(OrgUnit.id, OrgUnit.name)
                .order_by(func.count(wo_sub.c.id).desc())
                .limit(10)
            ).all()
            plant_bars = [{"name": str(n), "count": int(c)} for n, c in plant_rows]
            ctr_rows = self._db.execute(
                select(Contractor.name, func.count(wo_sub.c.id))
                .select_from(wo_sub)
                .join(Contractor, Contractor.id == wo_sub.c.contractor_id)
                .group_by(Contractor.id, Contractor.name)
                .order_by(func.count(wo_sub.c.id).desc())
                .limit(10)
            ).all()
            ctr_bars = [{"name": str(n), "count": int(c)} for n, c in ctr_rows]
            # aging pending (respects dashboard filters)
            today_d = date.today()
            aging = {"0-7": 0, "7-15": 0, "15-30": 0, "30+": 0}
            aging_q = select(WorkOrder).where(
                WorkOrder.status.in_(("draft", "pending_approval")),
                WorkOrder.id.in_(select(wo_sub.c.id)),
            )
            for w in self._db.scalars(aging_q).all():
                cd = w.created_at.date() if w.created_at else today_d
                days = (today_d - cd).days
                if days <= 7:
                    aging["0-7"] += 1
                elif days <= 15:
                    aging["7-15"] += 1
                elif days <= 30:
                    aging["15-30"] += 1
                else:
                    aging["30+"] += 1
            op_insights: list[str] = []
            if plant_bars:
                op_insights.append(f"Highest workload plant (current filters): {plant_bars[0]['name']}.")
            pend_ctr_rows = self._db.execute(
                select(Contractor.name, func.count(wo_sub.c.id))
                .select_from(wo_sub)
                .join(Contractor, Contractor.id == wo_sub.c.contractor_id)
                .where(wo_sub.c.status.in_(("draft", "pending_approval")))
                .group_by(Contractor.id, Contractor.name)
                .order_by(func.count(wo_sub.c.id).desc())
                .limit(1)
            ).first()
            if pend_ctr_rows:
                op_insights.append(f"Most pending work orders by contractor: {pend_ctr_rows[0]}.")
            if aging["30+"] > 0:
                op_insights.append(f"Delayed WOs: {aging['30+']} pending over 30 days.")
            if len(plant_bars) >= 2 and plant_bars[0]["count"] > 3 * max(1, plant_bars[1]["count"]):
                op_insights.append("Workload imbalance: top plant volume is more than 3× the next plant.")
            out["operations"] = {
                "donut_status": donut,
                "monthly_trend": wo_trend,
                "plant_workload": plant_bars,
                "contractor_allocation": ctr_bars,
                "pending_aging": [{"bucket": k, "count": v} for k, v in aging.items()],
                "insights": op_insights[:6],
            }

        # --- Negotiations / commercial ---
        if has("contractor_rates.view"):
            r_sub = self._rates_base(flt).subquery()
            by_status_rows = self._db.execute(
                select(r_sub.c.status, func.count()).select_from(r_sub).group_by(r_sub.c.status)
            ).all()
            funnel = {str(s or "unknown"): int(c) for s, c in by_status_rows}
            trend_m = []
            for y, m in _iter_last_n_months(8):
                s, e = _month_range(y, m)
                rc = select(func.count()).select_from(ContractorRate).where(
                    ContractorRate.created_at >= s,
                    ContractorRate.created_at < e,
                )
                if flt.contractor_id:
                    rc = rc.where(ContractorRate.contractor_id == int(flt.contractor_id))
                if flt.negotiation_status:
                    rc = rc.where(ContractorRate.status == str(flt.negotiation_status).strip())
                cnt = int(self._db.scalar(rc) or 0)
                trend_m.append({"month": f"{y}-{m:02d}", "count": cnt})
            # top premium contractors (sum max per contractor simplified: sum positive diff)
            prem_rows: list[tuple[str, float]] = []
            for cid, name, total_p in self._db.execute(
                select(Contractor.id, Contractor.name, func.sum(ContractorRate.negotiated_rate - PartMaster.base_rate))
                .join(ContractorRate, ContractorRate.contractor_id == Contractor.id)
                .join(PartMaster, PartMaster.id == ContractorRate.part_master_id)
                .where(ContractorRate.status == "approved", ContractorRate.negotiated_rate > PartMaster.base_rate)
                .group_by(Contractor.id, Contractor.name)
                .order_by(func.sum(ContractorRate.negotiated_rate - PartMaster.base_rate).desc())
                .limit(10)
            ).all():
                prem_rows.append((str(name), float(total_p or 0)))
            sav_rows = list(
                self._db.execute(
                    select(Contractor.name, func.coalesce(func.sum(ContractorRate.savings_amount), 0))
                    .join(ContractorRate, ContractorRate.contractor_id == Contractor.id)
                    .where(ContractorRate.status == "approved")
                    .group_by(Contractor.id, Contractor.name)
                    .order_by(func.coalesce(func.sum(ContractorRate.savings_amount), 0).desc())
                    .limit(10)
                ).all()
            )
            part_neg = list(
                self._db.execute(
                    select(PartMaster.part_code, func.count(ContractorRate.id))
                    .join(ContractorRate, ContractorRate.part_master_id == PartMaster.id)
                    .group_by(PartMaster.id, PartMaster.part_code)
                    .order_by(func.count(ContractorRate.id).desc())
                    .limit(10)
                ).all()
            )
            part_sav_rows = list(
                self._db.execute(
                    select(PartMaster.part_code, func.coalesce(func.sum(ContractorRate.savings_amount), 0))
                    .join(ContractorRate, ContractorRate.part_master_id == PartMaster.id)
                    .where(ContractorRate.status == "approved", ContractorRate.savings_amount.isnot(None))
                    .group_by(PartMaster.id, PartMaster.part_code)
                    .order_by(func.coalesce(func.sum(ContractorRate.savings_amount), 0).desc())
                    .limit(10)
                ).all()
            )
            savings_vs_premium: list[dict[str, Any]] = []
            for y, m in _iter_last_n_months(8):
                s, e = _month_range(y, m)
                sm = savings_month(s, e)
                prem_q = (
                    select(func.coalesce(func.sum(ContractorRate.negotiated_rate - PartMaster.base_rate), 0))
                    .join(PartMaster, PartMaster.id == ContractorRate.part_master_id)
                    .where(
                        ContractorRate.status == "approved",
                        ContractorRate.approved_at.isnot(None),
                        ContractorRate.approved_at >= s,
                        ContractorRate.approved_at < e,
                        ContractorRate.negotiated_rate > PartMaster.base_rate,
                    )
                )
                if flt.contractor_id:
                    prem_q = prem_q.where(ContractorRate.contractor_id == int(flt.contractor_id))
                pm = float(self._db.scalar(prem_q) or 0)
                savings_vs_premium.append({"month": f"{y}-{m:02d}", "savings": sm, "premium": pm})
            insights: list[str] = []
            if funnel.get("pending_approval", 0) + funnel.get("draft", 0) > 0:
                insights.append(
                    f"Pending negotiations: {funnel.get('pending_approval', 0) + funnel.get('draft', 0)} records awaiting workflow."
                )
            if prem_rows:
                insights.append(f"Highest premium contractor (sum above base): {prem_rows[0][0]}.")
            if sav_rows and float(sav_rows[0][1] or 0) > 0:
                insights.append(f"Largest recorded savings contributor: {sav_rows[0][0]}.")
            out["negotiations"] = {
                "funnel": funnel,
                "trend": trend_m,
                "savings_vs_premium": savings_vs_premium,
                "top_premium_contractors": [{"name": n, "premium": p} for n, p in prem_rows],
                "top_savings_contractors": [{"name": str(n), "savings": float(s or 0)} for n, s in sav_rows],
                "top_negotiated_parts": [{"part_code": str(pc), "count": int(c)} for pc, c in part_neg],
                "top_savings_parts": [{"part_code": str(pc), "savings": float(s or 0)} for pc, s in part_sav_rows],
                "insights": insights[:6],
            }

        # --- Invoices / financial ---
        if has("invoices.view", "invoices.create"):
            invs = list(self._db.scalars(self._inv_base(flt)).all())
            st_amt: dict[str, float] = {}
            for i in invs:
                st = str(i.status).lower()
                st_amt[st] = st_amt.get(st, 0.0) + float(_d0(i.total_amount))
            inv_donut = [{"name": k, "value": v} for k, v in sorted(st_amt.items(), key=lambda x: -x[1])]
            inv_sq = self._inv_base(flt).subquery()
            plant_spend = list(
                self._db.execute(
                    select(OrgUnit.name, func.coalesce(func.sum(inv_sq.c.total_amount), 0))
                    .select_from(inv_sq)
                    .join(OrgUnit, OrgUnit.id == inv_sq.c.org_unit_id)
                    .where(inv_sq.c.status.in_(("approved", "paid")))
                    .group_by(OrgUnit.id, OrgUnit.name)
                    .order_by(func.coalesce(func.sum(inv_sq.c.total_amount), 0).desc())
                    .limit(10)
                ).all()
            )
            ctr_spend = list(
                self._db.execute(
                    select(Contractor.name, func.coalesce(func.sum(inv_sq.c.total_amount), 0))
                    .select_from(inv_sq)
                    .join(Contractor, Contractor.id == inv_sq.c.contractor_id)
                    .where(inv_sq.c.status.in_(("approved", "paid")))
                    .group_by(Contractor.id, Contractor.name)
                    .order_by(func.coalesce(func.sum(inv_sq.c.total_amount), 0).desc())
                    .limit(10)
                ).all()
            )
            spend_trend = []
            for y, m in _iter_last_n_months(8):
                s, e = _month_range(y, m)
                amt_q = select(func.coalesce(func.sum(Invoice.total_amount), 0)).where(
                    Invoice.invoice_date >= s.date(),
                    Invoice.invoice_date < e.date(),
                    Invoice.status.in_(("approved", "paid")),
                )
                if flt.plant_id:
                    amt_q = amt_q.where(Invoice.org_unit_id == int(flt.plant_id))
                if flt.contractor_id:
                    amt_q = amt_q.where(Invoice.contractor_id == int(flt.contractor_id))
                amt = float(self._db.scalar(amt_q) or 0)
                spend_trend.append({"month": f"{y}-{m:02d}", "value": amt})
            inv_aging = {"0-3": 0, "4-7": 0, "8-14": 0, "15+": 0}
            pend_inv_st = ("submitted", "blocked", "pending_exception_approval")
            for i in invs:
                if str(i.status).lower() not in pend_inv_st:
                    continue
                anchor = i.submitted_at or i.created_at
                if not anchor:
                    continue
                days = (date.today() - anchor.date()).days
                if days <= 3:
                    inv_aging["0-3"] += 1
                elif days <= 7:
                    inv_aging["4-7"] += 1
                elif days <= 14:
                    inv_aging["8-14"] += 1
                else:
                    inv_aging["15+"] += 1
            fin_alerts: list[str] = []
            if exec_block.get("invoice_pending_value", 0) > 0 and exec_block.get("invoice_pending_value", 0) > exec_block.get(
                "invoice_approved_value", 0
            ):
                fin_alerts.append("Pending invoice exposure exceeds approved book (filtered view).")
            if ctr_spend and plant_spend:
                top_ctr = float(ctr_spend[0][1] or 0)
                second_ctr = float(ctr_spend[1][1] or 0) if len(ctr_spend) > 1 else 0.0
                if second_ctr > 0 and top_ctr > 4 * second_ctr:
                    fin_alerts.append(f"Spend concentration: {ctr_spend[0][0]} dominates approved spend in this slice.")
            rej_amt = st_amt.get("rejected", 0.0)
            if rej_amt > 0 and rej_amt > sum(st_amt.values()) * 0.15:
                fin_alerts.append("Rejected invoices represent a large share of filtered invoice value.")
            savings_vs_spend_trend: list[dict[str, Any]] = []
            for y, m in _iter_last_n_months(8):
                s, e = _month_range(y, m)
                spend_q = select(func.coalesce(func.sum(Invoice.total_amount), 0)).where(
                    Invoice.invoice_date >= s.date(),
                    Invoice.invoice_date < e.date(),
                    Invoice.status.in_(("approved", "paid")),
                )
                if flt.plant_id:
                    spend_q = spend_q.where(Invoice.org_unit_id == int(flt.plant_id))
                if flt.contractor_id:
                    spend_q = spend_q.where(Invoice.contractor_id == int(flt.contractor_id))
                spend_m = float(self._db.scalar(spend_q) or 0)
                sav_m = savings_month(s, e)
                savings_vs_spend_trend.append({"month": f"{y}-{m:02d}", "spend": spend_m, "savings": sav_m})
            out["financial"] = {
                "status_amounts": inv_donut,
                "plant_spend": [{"name": str(n), "value": float(v)} for n, v in plant_spend],
                "contractor_spend": [{"name": str(n), "value": float(v)} for n, v in ctr_spend],
                "monthly_spend_trend": spend_trend,
                "savings_vs_spend_trend": savings_vs_spend_trend,
                "invoice_approval_aging": [{"bucket": k, "count": v} for k, v in inv_aging.items()],
                "alerts": fin_alerts[:6],
            }

        # --- Parts (usage from WO lines) ---
        if has("work_orders.view", "work_orders.create"):
            wo_sq2 = self._wo_base(flt).subquery()
            pq = (
                select(PartMaster.part_code, func.count(WorkOrderItem.id), func.coalesce(func.sum(WorkOrderItem.taxable_value), 0))
                .select_from(WorkOrderItem)
                .join(wo_sq2, wo_sq2.c.id == WorkOrderItem.work_order_id)
                .join(PartMaster, PartMaster.id == WorkOrderItem.part_master_id)
            )
            if flt.part_pricing_method:
                pq = pq.where(PartMaster.pricing_method == str(flt.part_pricing_method).strip())
            part_rows = list(
                self._db.execute(
                    pq.group_by(PartMaster.id, PartMaster.part_code)
                    .order_by(func.count(WorkOrderItem.id).desc())
                    .limit(12)
                ).all()
            )
            out["parts"] = {
                "top_by_lines": [
                    {"part_code": str(pc), "lines": int(n), "value": float(v)} for pc, n, v in part_rows
                ]
            }

        # --- Pending actions (light) ---
        pending: list[dict[str, Any]] = []
        if has("contractor_rates.view"):
            for r in self._db.scalars(
                self._rates_base(flt).where(ContractorRate.status.in_(("draft", "pending_approval"))).limit(8)
            ).all():
                pending.append(
                    {
                        "kind": "negotiation",
                        "severity": "warning",
                        "title": f"Negotiation {r.status}",
                        "detail": f"Part master #{r.part_master_id}",
                        "href": f"/dashboard/negotiated-rates/{r.id}",
                    }
                )
        if has("work_orders.view"):
            for w in self._db.scalars(
                self._wo_base(flt)
                .where(WorkOrder.status.in_(("draft", "pending_approval")))
                .order_by(WorkOrder.created_at.desc())
                .limit(8)
            ).all():
                pending.append(
                    {
                        "kind": "work_order",
                        "severity": "warning",
                        "title": f"WO {w.work_order_number}",
                        "detail": w.title,
                        "href": f"/dashboard/work-orders/{w.id}",
                    }
                )
        if has("invoices.view"):
            for inv in self._db.scalars(
                self._inv_base(flt)
                .where(Invoice.status.in_(("draft", "submitted", "blocked", "pending_exception_approval")))
                .order_by(Invoice.created_at.desc())
                .limit(8)
            ).all():
                pending.append(
                    {
                        "kind": "invoice",
                        "severity": "danger" if inv.status == "blocked" else "warning",
                        "title": f"Invoice {inv.invoice_number}",
                        "detail": inv.status,
                        "href": f"/dashboard/invoices/{inv.id}",
                    }
                )
        out["pending_actions"] = pending[:24]

        if has("contractor.view") and (
            has("work_orders.view") or has("contractor_rates.view") or has("invoices.view")
        ):
            ops = out.get("operations") or {}
            neg = out.get("negotiations") or {}
            fin = out.get("financial") or {}
            out["contractor_intel"] = {
                "top_work_orders": (ops.get("contractor_allocation") or [])[:8],
                "top_premium": (neg.get("top_premium_contractors") or [])[:8],
                "top_savings": (neg.get("top_savings_contractors") or [])[:8],
                "top_invoice_spend": (fin.get("contractor_spend") or [])[:8],
            }

        # --- Timeline (recent) ---
        events: list[dict[str, Any]] = []
        if has("work_orders.view"):
            for w in self._db.scalars(select(WorkOrder).order_by(WorkOrder.created_at.desc()).limit(12)).all():
                events.append(
                    {
                        "at": w.created_at.isoformat() if w.created_at else "",
                        "type": "work_order",
                        "title": f"Work order {w.work_order_number}",
                        "subtitle": w.title,
                    }
                )
        if has("contractor_rates.view"):
            for r in self._db.scalars(select(ContractorRate).order_by(ContractorRate.updated_at.desc()).limit(12)).all():
                events.append(
                    {
                        "at": r.updated_at.isoformat() if r.updated_at else "",
                        "type": "negotiation",
                        "title": f"Negotiation updated ({r.status})",
                        "subtitle": f"Rate #{r.id}",
                    }
                )
        if has("invoices.view"):
            for inv in self._db.scalars(select(Invoice).order_by(Invoice.created_at.desc()).limit(12)).all():
                events.append(
                    {
                        "at": inv.created_at.isoformat() if inv.created_at else "",
                        "type": "invoice",
                        "title": f"Invoice {inv.invoice_number}",
                        "subtitle": inv.status,
                    }
                )
        events.sort(key=lambda e: e.get("at") or "", reverse=True)
        out["timeline"] = events[:30]

        return out

    def build_xlsx(self, grants: set[str], flt: IntelligenceFilters) -> bytes:
        try:
            from openpyxl import Workbook
        except ImportError as e:  # pragma: no cover
            raise RuntimeError("openpyxl required") from e

        data = self.build(grants, flt)
        wb = Workbook()
        ws = wb.active
        ws.title = "Executive"
        ws.append(["metric", "value"])
        ex = data.get("executive") or {}
        for k, v in ex.items():
            if k != "sparklines":
                ws.append([k, str(v)])

        if "operations" in data:
            w = wb.create_sheet("Work orders")
            for row in (data["operations"].get("monthly_trend") or []):
                w.append([row.get("month"), row.get("created"), row.get("completed")])

        if "negotiations" in data:
            n = wb.create_sheet("Negotiations")
            for row in (data["negotiations"].get("trend") or []):
                n.append([row.get("month"), row.get("count")])

        if "financial" in data:
            f = wb.create_sheet("Invoices")
            for row in (data["financial"].get("monthly_spend_trend") or []):
                f.append([row.get("month"), row.get("value")])

        if "financial" in data:
            fs = wb.create_sheet("Financial Summary")
            fs.append(["metric", "value"])
            fs.append(["savings_vs_spend rows", ""])
            for row in data["financial"].get("savings_vs_spend_trend") or []:
                fs.append([row.get("month"), row.get("spend"), row.get("savings")])

        if "operations" in data:
            pl = wb.create_sheet("Plant Analytics")
            pl.append(["plant", "work_orders"])
            for row in (data["operations"].get("plant_workload") or []):
                pl.append([row.get("name"), row.get("count")])

        if "contractor_intel" in data:
            ci = wb.create_sheet("Contractors")
            ci.append(["slice", "name", "metric"])
            for row in (data["contractor_intel"].get("top_work_orders") or []):
                ci.append(["work_orders", row.get("name"), row.get("count")])
            for row in (data["contractor_intel"].get("top_invoice_spend") or []):
                ci.append(["invoice_spend", row.get("name"), row.get("value")])

        pa = wb.create_sheet("Pending Actions")
        for row in data.get("pending_actions") or []:
            pa.append([row.get("kind"), row.get("severity"), row.get("title"), row.get("detail")])

        if "parts" in data:
            ps = wb.create_sheet("Parts")
            for row in (data["parts"].get("top_by_lines") or []):
                ps.append([row.get("part_code"), row.get("lines"), row.get("value")])

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
