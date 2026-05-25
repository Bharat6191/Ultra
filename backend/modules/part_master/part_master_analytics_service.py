"""Aggregated analytics for a single Part Master row (intelligence dashboard + export)."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from typing import Any, Sequence

from sqlalchemy import and_, func, select
from sqlalchemy.orm import Session

from modules.contractor.analytics_schema import AnalyticsTimelineEvent, NegotiationTrendPoint, PendingItem, PendingActions
from modules.contractor.models import Contractor
from modules.contractor_rates.models import ContractorRate, NegotiationLog
from modules.errors import NotFoundError
from modules.invoices.models import Invoice, InvoiceLine
from modules.org_units.model import OrgUnit
from modules.part_master.models import PartMaster, PartMasterAuditLog
from modules.users.model import User
from modules.work_orders.models import WorkOrder, WorkOrderItem, WorkOrderItemProgress

from modules.part_master.part_master_analytics_schema import (
    KpiCard,
    PartAnalyticsFilters,
    PartCommercialInsights,
    PartCompetitionAnalytics,
    PartCompetitionRow,
    PartContractorComparisonRow,
    PartContractorsAnalytics,
    PartMasterAnalyticsSummary,
    PartMasterHeaderBlock,
    PartNegotiationBundle,
    PartWorkOrderAnalytics,
    PartWorkOrderMonthly,
    PartWorkOrderRow,
)


def _d0(v: Decimal | None) -> Decimal:
    return v if v is not None else Decimal("0")


def _today() -> date:
    return datetime.now(timezone.utc).date()


class PartMasterAnalyticsService:
    def __init__(self, db: Session) -> None:
        self._db = db

    def _ensure(self, part_master_id: int) -> PartMaster:
        row = self._db.get(PartMaster, int(part_master_id))
        if row is None:
            raise NotFoundError("PartMaster", part_master_id)
        return row

    def _wo_date_cond(self, flt: PartAnalyticsFilters) -> list[Any]:
        cond: list[Any] = []
        if flt.date_from:
            cond.append(
                WorkOrder.created_at
                >= datetime.combine(flt.date_from, datetime.min.time()).replace(tzinfo=timezone.utc)
            )
        if flt.date_to:
            end = datetime.combine(flt.date_to + timedelta(days=1), datetime.min.time()).replace(tzinfo=timezone.utc)
            cond.append(WorkOrder.created_at < end)
        if flt.plant_id:
            cond.append(WorkOrder.org_unit_id == int(flt.plant_id))
        if flt.contractor_id:
            cond.append(WorkOrder.contractor_id == int(flt.contractor_id))
        if flt.work_order_status and str(flt.work_order_status).strip():
            cond.append(WorkOrder.status == str(flt.work_order_status).strip().lower())
        return cond

    def _rates_for_part(self, part_master_id: int, flt: PartAnalyticsFilters):
        stmt = select(ContractorRate).where(ContractorRate.part_master_id == int(part_master_id))
        if flt.negotiation_status and str(flt.negotiation_status).strip():
            stmt = stmt.where(ContractorRate.status == str(flt.negotiation_status).strip().lower())
        if flt.contractor_id:
            stmt = stmt.where(ContractorRate.contractor_id == int(flt.contractor_id))
        return list(self._db.scalars(stmt.order_by(ContractorRate.updated_at.desc())).all())

    def _wos_touching_part(self, part_master_id: int, flt: PartAnalyticsFilters) -> list[WorkOrder]:
        stmt = (
            select(WorkOrder)
            .join(WorkOrderItem, WorkOrderItem.work_order_id == WorkOrder.id)
            .where(WorkOrderItem.part_master_id == int(part_master_id))
            .distinct()
        )
        wc = self._wo_date_cond(flt)
        if wc:
            stmt = stmt.where(and_(*wc))
        return list(self._db.scalars(stmt).all())

    def _user_names(self, ids: set[int]) -> dict[int, str]:
        if not ids:
            return {}
        rows = list(self._db.scalars(select(User).where(User.id.in_(tuple(ids)))).all())
        return {int(u.id): str(u.full_name) for u in rows}

    def _contractor_wo_stats_for_part(self, part_master_id: int, flt: PartAnalyticsFilters) -> tuple[dict[int, int], dict[int, Decimal]]:
        """Per contractor: active WO count (approved+active with this part), sum taxable_value for this part."""
        wos = self._wos_touching_part(part_master_id, flt)
        wo_ids = [int(w.id) for w in wos]
        if not wo_ids:
            return {}, {}
        items = list(
            self._db.execute(
                select(WorkOrderItem, WorkOrder)
                .join(WorkOrder, WorkOrder.id == WorkOrderItem.work_order_id)
                .where(
                    WorkOrderItem.part_master_id == int(part_master_id),
                    WorkOrderItem.work_order_id.in_(wo_ids),
                )
            ).all()
        )
        active_ct: dict[int, int] = defaultdict(int)
        value_sum: dict[int, Decimal] = defaultdict(lambda: Decimal("0"))
        seen_pair: set[tuple[int, int]] = set()
        for it, wo in items:
            cid = int(wo.contractor_id)
            val = _d0(it.taxable_value)
            value_sum[cid] += val
            st = str(wo.status).lower()
            if st in ("approved", "active"):
                key = (cid, int(wo.id))
                if key not in seen_pair:
                    seen_pair.add(key)
                    active_ct[cid] += 1
        return dict(active_ct), dict(value_sum)

    def _item_completion_pct(self, item: WorkOrderItem) -> Decimal | None:
        latest: WorkOrderItemProgress | None = None
        if item.progress:
            latest = max(
                item.progress,
                key=lambda p: (
                    p.created_at or datetime.min.replace(tzinfo=timezone.utc),
                    int(p.id),
                ),
            )
        if latest is not None and latest.completed_percentage is not None:
            return Decimal(str(latest.completed_percentage))
        if (
            latest is not None
            and latest.completed_quantity is not None
            and item.planned_quantity is not None
        ):
            pq = Decimal(str(item.planned_quantity))
            if pq > 0:
                return Decimal(str(latest.completed_quantity)) / pq * Decimal("100")
        return None

    def _wo_completion_for_part(self, items: list[WorkOrderItem]) -> Decimal | None:
        pcts = [p for p in (self._item_completion_pct(it) for it in items) if p is not None]
        if not pcts:
            return None
        return sum(pcts) / Decimal(len(pcts))

    def get_summary(self, part_master_id: int, flt: PartAnalyticsFilters) -> PartMasterAnalyticsSummary:
        pm = self._ensure(part_master_id)
        rates = self._rates_for_part(part_master_id, flt)
        wos = self._wos_touching_part(part_master_id, flt)

        plant_ids = {int(w.org_unit_id) for w in wos}
        plant_names: list[str] = []
        plant_opts: list[dict[str, Any]] = []
        for pid in sorted(plant_ids):
            ou = self._db.get(OrgUnit, pid)
            nm = str(ou.name) if ou else str(pid)
            plant_names.append(nm)
            plant_opts.append({"id": pid, "name": nm})
        if not plant_opts:
            ou0 = self._db.get(OrgUnit, int(pm.org_unit_id))
            nm0 = str(ou0.name) if ou0 else f"Plant #{pm.org_unit_id}"
            plant_names = [nm0]
            plant_opts = [{"id": int(pm.org_unit_id), "name": nm0}]

        home = self._db.get(OrgUnit, int(pm.org_unit_id))
        contractors_from_rates = {int(r.contractor_id) for r in rates}
        contractors_from_wo = {int(w.contractor_id) for w in wos}
        all_c = contractors_from_rates | contractors_from_wo

        active_c = {
            int(w.contractor_id)
            for w in wos
            if str(w.status).lower() in ("approved", "active", "closed")
        } | {int(r.contractor_id) for r in rates if str(r.status).lower() == "approved"}

        neg_total = len(rates)
        neg_appr = sum(1 for r in rates if str(r.status).lower() == "approved")
        neg_pend = sum(1 for r in rates if str(r.status).lower() in ("draft", "pending_approval"))

        appr_negs = [Decimal(r.negotiated_rate) for r in rates if str(r.status).lower() == "approved"]
        low_n = min(appr_negs) if appr_negs else None
        high_n = max(appr_negs) if appr_negs else None
        avg_n = (sum(appr_negs) / len(appr_negs)) if appr_negs else None

        wo_total = len(wos)
        wo_active = sum(1 for w in wos if str(w.status).lower() in ("approved", "active"))
        wo_pend = sum(1 for w in wos if str(w.status).lower() in ("draft", "pending_approval"))

        if wos:
            wo_ids = [int(w.id) for w in wos]
            items = list(
                self._db.scalars(
                    select(WorkOrderItem).where(
                        WorkOrderItem.part_master_id == int(part_master_id),
                        WorkOrderItem.work_order_id.in_(wo_ids),
                    )
                ).all()
            )
        else:
            items = list(
                self._db.scalars(select(WorkOrderItem).where(WorkOrderItem.part_master_id == int(part_master_id))).all()
            )
        total_qty = sum(_d0(it.planned_quantity) for it in items)
        total_val = sum(_d0(it.taxable_value) for it in items)

        base = Decimal(pm.base_rate)
        appr_rates = [r for r in rates if str(r.status).lower() == "approved"]
        highest_prem_c = None
        best_sav_c = None
        if appr_rates:
            def prem(r: ContractorRate) -> Decimal:
                n = Decimal(r.negotiated_rate)
                return n - base if n > base else Decimal("0")

            hp = max(appr_rates, key=prem)
            if prem(hp) > 0:
                c = self._db.get(Contractor, int(hp.contractor_id))
                highest_prem_c = str(c.name) if c else None
            bs = max(appr_rates, key=lambda r: _d0(r.savings_amount))
            if _d0(bs.savings_amount) > 0:
                c2 = self._db.get(Contractor, int(bs.contractor_id))
                best_sav_c = str(c2.name) if c2 else None

        kpis: list[KpiCard] = [
            KpiCard(key="contractors_total", label="Total contractors", value=len(all_c), tone="neutral"),
            KpiCard(key="contractors_active", label="Active contractors", value=len(active_c), tone="neutral"),
            KpiCard(key="neg_total", label="Total negotiations", value=neg_total, tone="neutral"),
            KpiCard(key="neg_approved", label="Approved negotiations", value=neg_appr, tone="success"),
            KpiCard(key="neg_pending", label="Pending negotiations", value=neg_pend, tone="warning" if neg_pend else "neutral"),
            KpiCard(key="wo_total", label="Total work orders", value=wo_total, tone="neutral"),
            KpiCard(key="wo_pending", label="Pending work orders", value=wo_pend, tone="warning" if wo_pend else "neutral"),
            KpiCard(key="qty_total", label="Total consumption qty (planned)", value=total_qty, tone="neutral"),
            KpiCard(key="wo_value", label="Total work order value", value=total_val, tone="neutral"),
            KpiCard(
                key="lowest_rate",
                label="Lowest contractor rate (approved)",
                value=low_n if low_n is not None else "—",
                tone="success",
            ),
            KpiCard(
                key="highest_premium",
                label="Highest premium charged (est.)",
                value=highest_prem_c or "—",
                tone="warning" if highest_prem_c else "neutral",
            ),
            KpiCard(
                key="avg_savings",
                label="Avg savings (approved rows)",
                value=(
                    sum(_d0(r.savings_amount) for r in appr_rates) / len(appr_rates)
                    if appr_rates
                    else Decimal("0")
                ),
                tone="success",
            ),
        ]

        header = PartMasterHeaderBlock(
            part_master_id=int(pm.id),
            part_code=str(pm.part_code),
            part_name=str(pm.part_name),
            description=pm.description,
            part_category=str(pm.pricing_method),
            part_type=str(pm.rate_unit_type),
            unit_of_measurement=str(pm.unit_type),
            weight_per_unit=pm.weight_per_piece,
            base_rate=base,
            is_active=bool(pm.is_active),
            record_status=str(pm.status),
            home_plant_name=str(home.name) if home else None,
            plant_names_used=plant_names,
            plant_filter_options=plant_opts,
            total_contractors_touching=len(all_c),
            active_contractors=len(active_c),
            total_negotiation_records=neg_total,
            total_active_work_orders=int(wo_active),
            lowest_negotiated_rate=low_n,
            highest_negotiated_rate=high_n,
            average_negotiated_rate=avg_n,
        )
        return PartMasterAnalyticsSummary(header=header, kpis=kpis)

    def get_contractors(self, part_master_id: int, flt: PartAnalyticsFilters) -> PartContractorsAnalytics:
        pm = self._ensure(part_master_id)
        base = Decimal(pm.base_rate)
        rates = self._rates_for_part(part_master_id, flt)
        active_ct, value_sum = self._contractor_wo_stats_for_part(part_master_id, flt)

        uids: set[int] = set()
        for r in rates:
            if r.approved_by:
                uids.add(int(r.approved_by))
        names = self._user_names(uids)

        rows: list[PartContractorComparisonRow] = []
        trend_map: dict[str, list[Decimal]] = defaultdict(list)

        for r in rates:
            c = self._db.get(Contractor, int(r.contractor_id))
            rnds = len(r.negotiation_logs or [])
            last_ts = r.updated_at
            if r.negotiation_logs:
                last_ts = max(nl.created_at for nl in r.negotiation_logs if nl.created_at)
            neg = Decimal(r.negotiated_rate)
            prem_pct = ((neg - base) / base * Decimal("100")) if base > 0 else Decimal("0")
            rows.append(
                PartContractorComparisonRow(
                    contractor_id=int(r.contractor_id),
                    contractor_name=str(c.name) if c else f"#{r.contractor_id}",
                    contractor_code=c.contractor_code if c else None,
                    contractor_rate_id=int(r.id),
                    base_rate=base,
                    initial_rate=r.initial_rate,
                    final_negotiated_rate=neg,
                    savings_pct=r.savings_percentage,
                    premium_above_base_pct=prem_pct,
                    effective_from=r.effective_from,
                    effective_to=r.effective_to,
                    status=str(r.status),
                    rounds=rnds,
                    last_negotiation_at=last_ts,
                    approved_by_name=names.get(int(r.approved_by)) if r.approved_by else None,
                    active_work_orders=int(active_ct.get(int(r.contractor_id), 0)),
                    total_work_order_value=value_sum.get(int(r.contractor_id), Decimal("0")),
                )
            )
            key = (r.approved_at or r.created_at).strftime("%Y-%m") if (r.approved_at or r.created_at) else "unknown"
            trend_map[key].append(neg)

        trend = [
            NegotiationTrendPoint(period=k, count=len(vals), avg_negotiated=(sum(vals) / len(vals)) if vals else None)
            for k, vals in sorted(trend_map.items())
        ][-24:]

        today = _today()
        bar_rows: list[dict[str, Any]] = []
        scatter: list[dict[str, Any]] = []
        heatmap: list[dict[str, Any]] = []
        current_by_c: dict[int, tuple[str, Decimal, Decimal]] = {}

        for cid in sorted(set(int(r.contractor_id) for r in rates) | set(active_ct.keys())):
            c = self._db.get(Contractor, int(cid))
            nm = str(c.name) if c else f"#{cid}"
            appr = [r for r in rates if int(r.contractor_id) == cid and str(r.status).lower() == "approved"]
            pick = None
            for r in appr:
                ok_from = r.effective_from <= today
                ok_to = r.effective_to is None or r.effective_to >= today
                if ok_from and ok_to:
                    if pick is None or r.effective_from > pick.effective_from:
                        pick = r
            if pick is None and appr:
                pick = max(appr, key=lambda x: x.effective_from)
            if pick:
                nrate = Decimal(pick.negotiated_rate)
                prem = nrate - base if nrate > base else Decimal("0")
                prem_pct = (prem / base * Decimal("100")) if base > 0 else Decimal("0")
                current_by_c[cid] = (nm, nrate, prem_pct)
                bar_rows.append({"name": nm, "negotiated_rate": float(nrate), "base_rate": float(base)})
                scatter.append(
                    {
                        "contractor_id": cid,
                        "name": nm,
                        "work_orders": int(active_ct.get(cid, 0)),
                        "negotiated_rate": float(nrate),
                    }
                )
                heatmap.append({"contractor": nm, "premium_pct": float(prem_pct), "negotiated_rate": float(nrate)})

        heatmap.sort(key=lambda x: -x["premium_pct"])

        lb_low = None
        if current_by_c:
            cid, tpl = min(current_by_c.items(), key=lambda kv: kv[1][1])
            nm, rate, _prem = tpl
            lb_low = {"contractor_id": cid, "name": nm, "rate": float(rate)}
        lb_used = None
        if active_ct:
            cid = max(active_ct, key=lambda k: active_ct[k])
            c = self._db.get(Contractor, int(cid))
            lb_used = {"contractor_id": cid, "name": str(c.name) if c else f"#{cid}", "work_orders": int(active_ct[cid])}
        lb_sav = None
        sav_rates = [r for r in rates if str(r.status).lower() == "approved" and _d0(r.savings_amount) > 0]
        if sav_rates:
            r0 = max(sav_rates, key=lambda r: _d0(r.savings_amount))
            c = self._db.get(Contractor, int(r0.contractor_id))
            lb_sav = {
                "contractor_id": int(r0.contractor_id),
                "name": str(c.name) if c else f"#{r0.contractor_id}",
                "savings_amount": float(_d0(r0.savings_amount)),
            }

        leaderboard = {"lowest_cost": lb_low, "most_used": lb_used, "highest_savings": lb_sav}
        return PartContractorsAnalytics(
            rows=rows,
            bar_by_contractor=bar_rows,
            negotiation_trend=trend,
            scatter_volume=scatter,
            premium_heatmap=heatmap,
            leaderboard=leaderboard,
        )

    def get_negotiations(self, part_master_id: int, flt: PartAnalyticsFilters) -> PartNegotiationBundle:
        self._ensure(part_master_id)
        rates = self._rates_for_part(part_master_id, flt)
        status_counts: dict[str, int] = defaultdict(int)
        trend_map: dict[str, list[Decimal]] = defaultdict(list)
        for r in rates:
            st = str(r.status).lower()
            status_counts[st] += 1
            neg = Decimal(r.negotiated_rate)
            key = (r.approved_at or r.created_at).strftime("%Y-%m") if (r.approved_at or r.created_at) else "unknown"
            trend_map[key].append(neg)
        trend = [
            NegotiationTrendPoint(period=k, count=len(vals), avg_negotiated=(sum(vals) / len(vals)) if vals else None)
            for k, vals in sorted(trend_map.items())
        ][-36:]
        return PartNegotiationBundle(status_counts=dict(status_counts), trend=trend)

    def get_work_orders(self, part_master_id: int, flt: PartAnalyticsFilters) -> PartWorkOrderAnalytics:
        self._ensure(part_master_id)
        wos = self._wos_touching_part(part_master_id, flt)
        totals: dict[str, int] = defaultdict(int)
        by_plant: dict[int, dict[str, Any]] = {}
        by_ctr: dict[int, dict[str, Any]] = {}
        by_month_c: dict[str, int] = defaultdict(int)
        by_month_x: dict[str, int] = defaultdict(int)
        by_month_qty: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
        value_trend: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
        qty_trend: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))

        wo_ids = [int(w.id) for w in wos]
        items_map: dict[int, list[WorkOrderItem]] = defaultdict(list)
        if wo_ids:
            all_items = list(
                self._db.scalars(
                    select(WorkOrderItem).where(
                        WorkOrderItem.part_master_id == int(part_master_id),
                        WorkOrderItem.work_order_id.in_(wo_ids),
                    )
                ).all()
            )
            for it in all_items:
                items_map[int(it.work_order_id)].append(it)

        invoiced = Decimal("0")
        pending_inv = Decimal("0")
        inv_by_wo: dict[int, tuple[Decimal, Decimal]] = defaultdict(lambda: (Decimal("0"), Decimal("0")))
        if wo_ids:
            inv_rows = list(
                self._db.execute(
                    select(
                        WorkOrderItem.work_order_id,
                        Invoice.status,
                        func.coalesce(func.sum(InvoiceLine.amount), 0),
                    )
                    .join(InvoiceLine, InvoiceLine.invoice_id == Invoice.id)
                    .join(WorkOrderItem, WorkOrderItem.id == InvoiceLine.work_order_item_id)
                    .where(
                        WorkOrderItem.part_master_id == int(part_master_id),
                        WorkOrderItem.work_order_id.in_(wo_ids),
                    )
                    .group_by(WorkOrderItem.work_order_id, Invoice.status)
                ).all()
            )
            for wo_id, st, amt in inv_rows:
                a = Decimal(str(amt))
                st_l = str(st).lower()
                cur_inv, cur_pend = inv_by_wo[int(wo_id)]
                if st_l in ("approved", "paid"):
                    invoiced += a
                    inv_by_wo[int(wo_id)] = (cur_inv + a, cur_pend)
                elif st_l in ("draft", "submitted", "blocked", "pending_exception_approval"):
                    pending_inv += a
                    inv_by_wo[int(wo_id)] = (cur_inv, cur_pend + a)

        total_qty = Decimal("0")
        wo_rows: list[PartWorkOrderRow] = []
        for w in wos:
            st = str(w.status).lower()
            totals[st] += 1
            oid = int(w.org_unit_id)
            if oid not in by_plant:
                ou = self._db.get(OrgUnit, oid)
                by_plant[oid] = {"plant_id": oid, "plant_name": ou.name if ou else str(oid), "count": 0, "value": Decimal("0")}
            by_plant[oid]["count"] += 1
            cid = int(w.contractor_id)
            if cid not in by_ctr:
                co = self._db.get(Contractor, cid)
                by_ctr[cid] = {
                    "contractor_id": cid,
                    "contractor_name": str(co.name) if co else str(cid),
                    "count": 0,
                    "value": Decimal("0"),
                }
            by_ctr[cid]["count"] += 1
            mkey = w.created_at.strftime("%Y-%m") if w.created_at else "unknown"
            by_month_c[mkey] += 1
            if st == "closed" and w.updated_at:
                by_month_x[w.updated_at.strftime("%Y-%m")] += 1
            line_val = Decimal("0")
            line_qty = Decimal("0")
            for it in items_map.get(int(w.id), []):
                line_val += _d0(it.taxable_value)
                pq = _d0(it.planned_quantity)
                line_qty += pq
                total_qty += pq
            by_plant[oid]["value"] += line_val
            by_ctr[cid]["value"] += line_val
            value_trend[mkey] += line_val
            qty_trend[mkey] += line_qty
            by_month_qty[mkey] += line_qty

            co = self._db.get(Contractor, int(w.contractor_id))
            wo_inv, wo_pend = inv_by_wo.get(int(w.id), (Decimal("0"), Decimal("0")))
            comp = self._wo_completion_for_part(items_map.get(int(w.id), []))
            st_l = str(w.status).lower()
            end_dt = w.updated_at.date() if st_l == "closed" and w.updated_at else None
            wo_rows.append(
                PartWorkOrderRow(
                    work_order_id=int(w.id),
                    work_order_number=str(w.work_order_number),
                    contractor_id=int(w.contractor_id),
                    contractor_name=str(co.name) if co else str(w.contractor_id),
                    quantity=line_qty,
                    wo_value=line_val,
                    invoiced_value=wo_inv,
                    pending_value=wo_pend,
                    status=st_l,
                    completion_pct=comp,
                    start_date=w.created_at.date() if w.created_at else None,
                    end_date=end_dt,
                )
            )

        wo_rows.sort(key=lambda r: r.work_order_number, reverse=True)

        months = sorted(set(by_month_c) | set(by_month_x) | set(value_trend.keys()) | set(qty_trend.keys()))
        tail = months[-24:]
        by_month_list = [
            PartWorkOrderMonthly(
                month=m,
                created=by_month_c.get(m, 0),
                completed=by_month_x.get(m, 0),
                quantity=by_month_qty.get(m, Decimal("0")),
            )
            for m in tail
        ]
        vtrend = [{"month": m, "value": float(value_trend.get(m, Decimal("0")))} for m in tail]
        qtrend = [{"month": m, "quantity": float(qty_trend.get(m, Decimal("0")))} for m in tail]
        donut = [{"name": k, "value": v} for k, v in sorted(totals.items(), key=lambda kv: -kv[1])]
        total_val = sum(float(x["value"]) for x in by_plant.values())
        avg_qty = (total_qty / len(wos)) if wos else None

        return PartWorkOrderAnalytics(
            totals_by_status=dict(totals),
            by_plant=[{**p, "value": float(p["value"])} for p in by_plant.values()],
            by_contractor=[{**c, "value": float(c["value"])} for c in by_ctr.values()],
            by_month=by_month_list,
            total_wo_value=Decimal(str(total_val)),
            total_invoiced_for_part=invoiced,
            pending_invoice_amount_for_part=pending_inv,
            donut_status=donut,
            value_trend=vtrend,
            qty_trend=qtrend,
            total_consumption_qty=total_qty,
            avg_order_quantity=avg_qty,
            rows=wo_rows,
        )

    def get_commercial(self, part_master_id: int, flt: PartAnalyticsFilters) -> PartCommercialInsights:
        pm = self._ensure(part_master_id)
        base = Decimal(pm.base_rate)
        rates = [r for r in self._rates_for_part(part_master_id, flt) if str(r.status).lower() == "approved"]
        if not rates:
            return PartCommercialInsights(
                lowest_negotiated_rate=None,
                highest_negotiated_rate=None,
                average_market_negotiated_rate=None,
                spread_low_high=None,
                total_savings_through_negotiation=Decimal("0"),
                total_extra_above_base=Decimal("0"),
                highest_premium_contractor_name=None,
                best_value_contractor_name=None,
                pct_rates_below_base=Decimal("0"),
                pct_rates_above_base=Decimal("0"),
                risk_level="low",
                insight_lines=["No approved negotiated rates for this part yet."],
                cost_trend=[],
            )
        negs = [Decimal(r.negotiated_rate) for r in rates]
        low = min(negs)
        high = max(negs)
        avg = sum(negs) / len(negs)
        extra = sum((n - base) for n in negs if n > base)
        total_sav = sum(_d0(r.savings_amount) for r in rates)
        below = sum(1 for n in negs if n < base)
        above = sum(1 for n in negs if n > base)
        n = len(negs)
        hp = max(rates, key=lambda r: max(Decimal(r.negotiated_rate) - base, Decimal("0")))
        c_hp = self._db.get(Contractor, int(hp.contractor_id))
        bs = max(rates, key=lambda r: _d0(r.savings_amount))
        c_bs = self._db.get(Contractor, int(bs.contractor_id))
        risk = "low"
        if high - low > base * Decimal("0.15") or above / n > Decimal("0.5"):
            risk = "high"
        elif high - low > base * Decimal("0.08") or above / n > Decimal("0.35"):
            risk = "medium"

        lines: list[str] = []
        if n > 1 and low > 0:
            lines.append(f"Approved negotiated rates span ₹{float(high - low):,.2f} between lowest and highest contractor.")
        if avg > base:
            lines.append(f"Average negotiated rate ({float(avg):,.2f}) is above the part base ({float(base):,.2f}).")
        elif avg < base:
            lines.append(f"Average negotiated rate ({float(avg):,.2f}) is below the part base ({float(base):,.2f}).")
        if c_hp and Decimal(hp.negotiated_rate) > base:
            prem_pct = (Decimal(hp.negotiated_rate) - base) / base * 100 if base > 0 else Decimal("0")
            lines.append(f"{c_hp.name} shows the highest premium ({float(prem_pct):.1f}% above base) among approved rows.")
        if c_bs and _d0(bs.savings_amount) > 0:
            lines.append(f"{c_bs.name} records the largest negotiation savings amount on this part.")

        trend_map: dict[str, list[Decimal]] = defaultdict(list)
        for r in rates:
            key = (r.approved_at or r.created_at).strftime("%Y-%m") if (r.approved_at or r.created_at) else "unknown"
            trend_map[key].append(Decimal(r.negotiated_rate))
        cost_trend = [
            {"month": k, "avg_rate": float(sum(v) / len(v))} for k, v in sorted(trend_map.items())
        ][-24:]

        return PartCommercialInsights(
            lowest_negotiated_rate=low,
            highest_negotiated_rate=high,
            average_market_negotiated_rate=avg,
            spread_low_high=high - low,
            total_savings_through_negotiation=total_sav,
            total_extra_above_base=extra,
            highest_premium_contractor_name=str(c_hp.name) if c_hp else None,
            best_value_contractor_name=str(c_bs.name) if c_bs else None,
            pct_rates_below_base=Decimal(100 * below / n) if n else Decimal("0"),
            pct_rates_above_base=Decimal(100 * above / n) if n else Decimal("0"),
            risk_level=risk,
            insight_lines=lines[:8],
            cost_trend=cost_trend,
        )

    def get_competition(self, part_master_id: int, flt: PartAnalyticsFilters) -> PartCompetitionAnalytics:
        self._ensure(part_master_id)
        rates = self._rates_for_part(part_master_id, flt)
        wos = self._wos_touching_part(part_master_id, flt)
        by_c: dict[int, list[WorkOrder]] = defaultdict(list)
        for w in wos:
            by_c[int(w.contractor_id)].append(w)

        contractors = sorted(set(int(r.contractor_id) for r in rates) | set(by_c.keys()))
        rate_rank: list[tuple[int, Decimal]] = []
        today = _today()
        for cid in contractors:
            appr = [r for r in rates if int(r.contractor_id) == cid and str(r.status).lower() == "approved"]
            pick = None
            for r in appr:
                if r.effective_from <= today and (r.effective_to is None or r.effective_to >= today):
                    if pick is None or r.effective_from > pick.effective_from:
                        pick = r
            if pick is None and appr:
                pick = max(appr, key=lambda x: x.effective_from)
            if pick:
                rate_rank.append((cid, Decimal(pick.negotiated_rate)))
            else:
                rate_rank.append((cid, Decimal("999999999")))

        rate_rank.sort(key=lambda x: x[1])
        rank_map = {cid: i + 1 for i, (cid, _) in enumerate(rate_rank)}

        rows: list[PartCompetitionRow] = []
        for cid in contractors:
            c = self._db.get(Contractor, int(cid))
            wl = by_c.get(cid, [])
            closed = sum(1 for w in wl if str(w.status).lower() == "closed")
            total = len(wl)
            comp_pct = (Decimal(100 * closed / total) if total else None)
            appr_n = sum(1 for r in rates if int(r.contractor_id) == cid and str(r.status).lower() == "approved")
            rej_n = sum(1 for r in rates if int(r.contractor_id) == cid and str(r.status).lower() == "rejected")
            denom = appr_n + rej_n
            succ_pct = (Decimal(100 * appr_n / denom) if denom else None)
            pend = sum(
                1 for r in rates if int(r.contractor_id) == cid and str(r.status).lower() in ("draft", "pending_approval")
            )
            crs = [r for r in rates if int(r.contractor_id) == cid]
            avg_r = (
                sum(len(r.negotiation_logs or []) for r in crs) / len(crs)
                if crs
                else None
            )
            rows.append(
                PartCompetitionRow(
                    contractor_id=cid,
                    contractor_name=str(c.name) if c else f"#{cid}",
                    rate_rank=int(rank_map.get(cid, 99)),
                    work_order_count=total,
                    completion_rate_pct=comp_pct,
                    negotiation_success_pct=succ_pct,
                    pending_approvals=pend,
                    avg_rounds=Decimal(str(round(avg_r, 2))) if avg_r is not None else None,
                )
            )
        rows.sort(key=lambda r: (r.rate_rank, r.work_order_count * -1))
        return PartCompetitionAnalytics(rows=rows)

    def get_pending(self, part_master_id: int) -> PendingActions:
        pm = self._ensure(part_master_id)
        items: list[PendingItem] = []
        today = _today()
        warn_until = today + timedelta(days=30)

        for r in self._db.scalars(
            select(ContractorRate).where(
                ContractorRate.part_master_id == int(part_master_id),
                ContractorRate.status.in_(("draft", "pending_approval")),
            )
        ).all():
            c = self._db.get(Contractor, int(r.contractor_id))
            items.append(
                PendingItem(
                    kind="negotiation",
                    severity="warning" if r.status == "pending_approval" else "info",
                    title=f"Negotiation {str(r.status).replace('_', ' ')} — {c.name if c else r.contractor_id}",
                    detail=f"Part {pm.part_code}",
                    href_hint=f"/dashboard/negotiated-rates/{r.id}",
                    entity_id=int(r.id),
                )
            )

        for r in self._db.scalars(
            select(ContractorRate).where(
                ContractorRate.part_master_id == int(part_master_id),
                ContractorRate.status == "approved",
                ContractorRate.effective_to.isnot(None),
                ContractorRate.effective_to >= today,
                ContractorRate.effective_to <= warn_until,
            )
        ).all():
            items.append(
                PendingItem(
                    kind="rate_expiry",
                    severity="warning",
                    title="Negotiated rate expiring soon",
                    detail=f"Effective to {r.effective_to}",
                    href_hint=f"/dashboard/negotiated-rates/{r.id}",
                    entity_id=int(r.id),
                )
            )

        base = Decimal(pm.base_rate)
        for r in self._db.scalars(
            select(ContractorRate).where(
                ContractorRate.part_master_id == int(part_master_id),
                ContractorRate.status == "approved",
            )
        ).all():
            n = Decimal(r.negotiated_rate)
            if base > 0 and n > base * Decimal("1.12"):
                c = self._db.get(Contractor, int(r.contractor_id))
                items.append(
                    PendingItem(
                        kind="premium",
                        severity="danger",
                        title=f"High premium vs base — {c.name if c else r.contractor_id}",
                        detail=f"Negotiated {n} vs base {base}",
                        href_hint=f"/dashboard/negotiated-rates/{r.id}",
                        entity_id=int(r.id),
                    )
                )

        for w in self._db.scalars(
            select(WorkOrder)
            .join(WorkOrderItem, WorkOrderItem.work_order_id == WorkOrder.id)
            .where(
                WorkOrderItem.part_master_id == int(part_master_id),
                WorkOrder.status.in_(("draft", "pending_approval", "active")),
            )
            .distinct()
        ).all():
            items.append(
                PendingItem(
                    kind="work_order",
                    severity="warning" if w.status != "active" else "info",
                    title=f"Work order {w.work_order_number} ({w.status})",
                    detail=w.title,
                    href_hint=f"/dashboard/work-orders/{w.id}",
                    entity_id=int(w.id),
                )
            )

        return PendingActions(items=items)

    def get_timeline(self, part_master_id: int, categories: Sequence[str] | None) -> list[AnalyticsTimelineEvent]:
        pm = self._ensure(part_master_id)
        cats = {str(c).lower() for c in categories} if categories else {"all"}
        use_all = "all" in cats
        out: list[AnalyticsTimelineEvent] = []

        def allow(cat: str) -> bool:
            return use_all or cat in cats

        if allow("negotiations") or allow("approvals"):
            for r in self._db.scalars(select(ContractorRate).where(ContractorRate.part_master_id == int(part_master_id))).all():
                c = self._db.get(Contractor, int(r.contractor_id))
                label = c.name if c else str(r.contractor_id)
                if allow("negotiations"):
                    out.append(
                        AnalyticsTimelineEvent(
                            category="negotiations",
                            type="rate_record",
                            title=f"Negotiation — {label}",
                            description=f"Status {r.status}",
                            timestamp=r.created_at,
                            metadata={"contractor_rate_id": int(r.id)},
                        )
                    )
                for log in r.negotiation_logs or []:
                    if allow("negotiations"):
                        out.append(
                            AnalyticsTimelineEvent(
                                category="negotiations",
                                type="negotiation_round",
                                title=f"Round {log.round_number} — {label}",
                                description=log.remarks,
                                timestamp=log.created_at,
                                metadata={"round": log.round_number},
                            )
                        )
                if allow("approvals") and r.approved_at:
                    out.append(
                        AnalyticsTimelineEvent(
                            category="approvals",
                            type="rate_approved",
                            title=f"Negotiation approved — {label}",
                            description=None,
                            timestamp=r.approved_at,
                            metadata={"contractor_rate_id": int(r.id)},
                        )
                    )

        if allow("work_orders"):
            for w in self._db.scalars(
                select(WorkOrder)
                .join(WorkOrderItem, WorkOrderItem.work_order_id == WorkOrder.id)
                .where(WorkOrderItem.part_master_id == int(part_master_id))
                .distinct()
            ).all():
                out.append(
                    AnalyticsTimelineEvent(
                        category="work_orders",
                        type="wo_created",
                        title=f"Work order {w.work_order_number}",
                        description=w.title,
                        timestamp=w.created_at,
                        metadata={"work_order_id": int(w.id)},
                    )
                )

        if allow("financial"):
            invs = list(
                self._db.scalars(
                    select(Invoice)
                    .join(InvoiceLine, InvoiceLine.invoice_id == Invoice.id)
                    .join(WorkOrderItem, WorkOrderItem.id == InvoiceLine.work_order_item_id)
                    .where(WorkOrderItem.part_master_id == int(part_master_id))
                    .distinct()
                ).all()
            )
            for inv in invs:
                out.append(
                    AnalyticsTimelineEvent(
                        category="financial",
                        type="invoice",
                        title=f"Invoice {inv.invoice_number} ({inv.status})",
                        description=None,
                        timestamp=inv.created_at,
                        metadata={"invoice_id": int(inv.id)},
                    )
                )

        if allow("compliance"):
            for a in self._db.scalars(
                select(PartMasterAuditLog)
                .where(PartMasterAuditLog.part_master_id == int(part_master_id))
                .order_by(PartMasterAuditLog.created_at.desc())
            ).all():
                out.append(
                    AnalyticsTimelineEvent(
                        category="compliance",
                        type="part_audit",
                        title=f"Part Master {a.action}",
                        description=None,
                        timestamp=a.created_at,
                        metadata={"audit_id": int(a.id)},
                    )
                )

        out.sort(key=lambda e: e.timestamp, reverse=True)
        return out[:500]

    def build_xlsx(self, part_master_id: int, flt: PartAnalyticsFilters) -> bytes:
        try:
            from openpyxl import Workbook
        except ImportError as e:  # pragma: no cover
            raise RuntimeError("openpyxl is required for Excel export. pip install openpyxl") from e

        self._ensure(part_master_id)
        wb = Workbook()
        ws0 = wb.active
        ws0.title = "Summary"
        summ = self.get_summary(part_master_id, flt)
        h = summ.header
        ws0.append(["field", "value"])
        for k, v in h.model_dump().items():
            ws0.append([k, str(v)])

        ws1 = wb.create_sheet("Contractor rates")
        ctr = self.get_contractors(part_master_id, flt)
        ws1.append(
            [
                "contractor",
                "rate_id",
                "status",
                "base",
                "initial",
                "negotiated",
                "savings_pct",
                "premium_pct",
                "effective_from",
                "effective_to",
                "rounds",
                "active_wo",
                "wo_value",
            ]
        )
        for r in ctr.rows:
            ws1.append(
                [
                    r.contractor_name,
                    r.contractor_rate_id,
                    r.status,
                    float(r.base_rate),
                    float(r.initial_rate) if r.initial_rate is not None else None,
                    float(r.final_negotiated_rate),
                    float(r.savings_pct) if r.savings_pct is not None else None,
                    float(r.premium_above_base_pct),
                    r.effective_from.isoformat(),
                    r.effective_to.isoformat() if r.effective_to else None,
                    r.rounds,
                    r.active_work_orders,
                    float(r.total_work_order_value),
                ]
            )

        ws2 = wb.create_sheet("Work orders")
        wo = self.get_work_orders(part_master_id, flt)
        ws2.append(["status", "count"])
        for k, v in sorted(wo.totals_by_status.items()):
            ws2.append([k, v])
        ws2.append([])
        ws2.append(["contractor", "count", "value"])
        for c in wo.by_contractor:
            ws2.append([c.get("contractor_name"), c.get("count"), c.get("value")])

        ws3 = wb.create_sheet("Commercial")
        com = self.get_commercial(part_master_id, flt)
        ws3.append(["metric", "value"])
        for k, v in com.model_dump().items():
            if k not in ("insight_lines", "cost_trend"):
                ws3.append([k, str(v)])
        ws3.append([])
        ws3.append(["insights"])
        for line in com.insight_lines:
            ws3.append([line])

        ws3b = wb.create_sheet("Competition")
        comp = self.get_competition(part_master_id, flt)
        ws3b.append(["rank", "contractor", "wo_count", "completion_pct", "success_pct", "pending_negs", "avg_rounds"])
        for r in comp.rows:
            ws3b.append(
                [
                    r.rate_rank,
                    r.contractor_name,
                    r.work_order_count,
                    float(r.completion_rate_pct) if r.completion_rate_pct is not None else None,
                    float(r.negotiation_success_pct) if r.negotiation_success_pct is not None else None,
                    r.pending_approvals,
                    float(r.avg_rounds) if r.avg_rounds is not None else None,
                ]
            )

        ws4 = wb.create_sheet("Pending")
        pend = self.get_pending(part_master_id)
        ws4.append(["kind", "severity", "title", "detail"])
        for it in pend.items:
            ws4.append([it.kind, it.severity, it.title, it.detail or ""])

        ws5 = wb.create_sheet("Timeline")
        tl = self.get_timeline(part_master_id, ("all",))
        ws5.append(["timestamp", "category", "type", "title"])
        for e in tl[:400]:
            ws5.append([e.timestamp.isoformat(), e.category, e.type, e.title])

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
