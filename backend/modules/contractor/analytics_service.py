"""Aggregated analytics for a single contractor (dashboard + exports)."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from typing import Any, Sequence

from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from modules.approvals.model import ApprovalRequest, ApprovalTask
from modules.contractor.models import Contractor, ContractorDocument, ContractorPlant
from modules.contractor_rates.models import ContractorRate, NegotiationLog
from modules.errors import NotFoundError
from modules.invoices.models import Invoice
from modules.org_units.model import OrgUnit
from modules.part_master.models import PartMaster
from modules.work_orders.models import WorkOrder, WorkOrderItem

from modules.contractor.analytics_schema import (
    AnalyticsFilters,
    AnalyticsTimelineEvent,
    CommercialInsights,
    ContractorAnalyticsSummary,
    KpiCard,
    NegotiationAnalytics,
    NegotiationRow,
    NegotiationTrendPoint,
    PendingActions,
    PendingItem,
    WorkOrderAnalytics,
    WorkOrderMonthly,
)


def _d0(v: Decimal | None) -> Decimal:
    return v if v is not None else Decimal("0")


def _today() -> date:
    return datetime.now(timezone.utc).date()


class ContractorAnalyticsService:
    def __init__(self, db: Session) -> None:
        self._db = db

    def _ensure(self, contractor_id: int) -> Contractor:
        row = self._db.get(Contractor, int(contractor_id))
        if row is None:
            raise NotFoundError("Contractor", contractor_id)
        return row

    def _part_filter(self, flt: AnalyticsFilters):
        cond = []
        if flt.part_search and str(flt.part_search).strip():
            q = f"%{str(flt.part_search).strip()}%"
            cond.append(or_(PartMaster.part_code.ilike(q), PartMaster.part_name.ilike(q)))
        return and_(*cond) if cond else None

    def _wo_date_filter(self, flt: AnalyticsFilters):
        cond = []
        if flt.date_from:
            cond.append(WorkOrder.created_at >= datetime.combine(flt.date_from, datetime.min.time()).replace(tzinfo=timezone.utc))
        if flt.date_to:
            end = datetime.combine(flt.date_to + timedelta(days=1), datetime.min.time()).replace(tzinfo=timezone.utc)
            cond.append(WorkOrder.created_at < end)
        if flt.plant_id:
            cond.append(WorkOrder.org_unit_id == int(flt.plant_id))
        if flt.work_order_status and str(flt.work_order_status).strip():
            cond.append(WorkOrder.status == str(flt.work_order_status).strip().lower())
        return and_(*cond) if cond else None

    def _rate_filters(self, flt: AnalyticsFilters):
        cond = []
        pf = self._part_filter(flt)
        if pf is not None:
            cond.append(pf)
        if flt.negotiation_status and str(flt.negotiation_status).strip():
            cond.append(ContractorRate.status == str(flt.negotiation_status).strip().lower())
        return and_(*cond) if cond else None

    def get_summary(self, contractor_id: int, flt: AnalyticsFilters) -> ContractorAnalyticsSummary:
        c = self._ensure(contractor_id)

        plants = list(
            self._db.execute(
                select(ContractorPlant.org_unit_id, OrgUnit.name)
                .join(OrgUnit, OrgUnit.id == ContractorPlant.org_unit_id)
                .where(ContractorPlant.contractor_id == int(contractor_id))
                .order_by(OrgUnit.name.asc())
            ).all()
        )
        plant_names = [str(r[1]) for r in plants]
        plant_opts = [{"id": int(r[0]), "name": str(r[1])} for r in plants]
        active_since = self._db.scalar(
            select(func.min(ContractorPlant.start_date)).where(
                ContractorPlant.contractor_id == int(contractor_id),
                ContractorPlant.start_date.isnot(None),
            )
        )
        if active_since is None and c.created_at:
            active_since = c.created_at.date() if hasattr(c.created_at, "date") else None

        wo_base = select(WorkOrder).where(WorkOrder.contractor_id == int(contractor_id))
        wox = self._wo_date_filter(flt)
        if wox is not None:
            wo_base = wo_base.where(wox)

        wo_rows = list(self._db.scalars(wo_base).all())
        by_st = defaultdict(int)
        for w in wo_rows:
            by_st[str(w.status).lower()] += 1

        total_wo = len(wo_rows)
        pending_wo = sum(1 for w in wo_rows if w.status in ("draft", "pending_approval"))
        in_prog = sum(1 for w in wo_rows if w.status == "active")
        completed = sum(1 for w in wo_rows if w.status == "closed")
        cancelled = sum(1 for w in wo_rows if w.status in ("cancelled", "rejected"))

        rate_stmt = (
            select(ContractorRate)
            .join(PartMaster, PartMaster.id == ContractorRate.part_master_id)
            .where(ContractorRate.contractor_id == int(contractor_id))
        )
        rf = self._rate_filters(flt)
        if rf is not None:
            rate_stmt = rate_stmt.where(rf)
        rates = list(self._db.scalars(rate_stmt).all())

        neg_total = len(rates)
        neg_approved = sum(1 for r in rates if str(r.status).lower() == "approved")
        neg_pending = sum(1 for r in rates if str(r.status).lower() in ("draft", "pending_approval"))
        neg_rejected = sum(1 for r in rates if str(r.status).lower() == "rejected")
        distinct_parts = len({int(r.part_master_id) for r in rates})
        active_wo_ct = sum(1 for w in wo_rows if w.status in ("approved", "active"))

        total_savings = sum((_d0(r.savings_amount) for r in rates if r.status == "approved"), start=Decimal("0"))
        premium = Decimal("0")
        for r in rates:
            if r.status != "approved":
                continue
            pm = self._db.get(PartMaster, int(r.part_master_id))
            if pm is None:
                continue
            diff = Decimal(r.negotiated_rate) - Decimal(pm.base_rate)
            if diff > 0:
                premium += diff

        kpis: list[KpiCard] = [
            KpiCard(key="wo_total", label="Total work orders", value=total_wo, tone="neutral"),
            KpiCard(key="wo_pending", label="Pending work orders", value=pending_wo, tone="warning" if pending_wo else "neutral"),
            KpiCard(key="wo_active", label="In progress (active)", value=in_prog, tone="neutral"),
            KpiCard(key="wo_completed", label="Completed (closed)", value=completed, tone="success"),
            KpiCard(key="wo_cancelled", label="Cancelled / rejected WO", value=cancelled, tone="neutral"),
            KpiCard(key="neg_total", label="Total negotiations", value=neg_total, tone="neutral"),
            KpiCard(key="neg_approved", label="Approved negotiations", value=neg_approved, tone="success"),
            KpiCard(key="neg_rejected", label="Rejected negotiations", value=neg_rejected, tone="danger" if neg_rejected else "neutral"),
            KpiCard(key="neg_pending", label="Pending / draft negotiations", value=neg_pending, tone="warning" if neg_pending else "neutral"),
            KpiCard(key="savings_total", label="Total negotiation savings", value=total_savings, tone="success"),
            KpiCard(key="premium_total", label="Premium above base (approved)", value=premium, tone="warning" if premium > 0 else "neutral"),
        ]

        return ContractorAnalyticsSummary(
            contractor_id=int(c.id),
            name=str(c.name),
            contractor_code=c.contractor_code,
            status=str(c.status),
            is_active=bool(c.is_active),
            plant_names=plant_names,
            plants=plant_opts,
            active_since=active_since if isinstance(active_since, date) else None,
            total_active_work_orders=int(active_wo_ct),
            distinct_negotiated_parts=int(distinct_parts),
            approved_negotiations=int(neg_approved),
            overall_savings_vs_initial=total_savings,
            overall_premium_above_base=premium,
            kpis=kpis,
        )

    def get_negotiations(self, contractor_id: int, flt: AnalyticsFilters) -> NegotiationAnalytics:
        self._ensure(contractor_id)
        stmt = (
            select(ContractorRate, PartMaster)
            .join(PartMaster, PartMaster.id == ContractorRate.part_master_id)
            .where(ContractorRate.contractor_id == int(contractor_id))
            .order_by(ContractorRate.updated_at.desc())
        )
        rf = self._rate_filters(flt)
        if rf is not None:
            stmt = stmt.where(rf)
        pairs = list(self._db.execute(stmt).all())

        rows: list[NegotiationRow] = []
        status_counts: dict[str, int] = defaultdict(int)
        bar_parts: list[dict[str, Any]] = []
        trend_map: dict[str, list[Decimal]] = defaultdict(list)
        premium_rank: list[tuple[str, Decimal]] = []

        for cr, pm in pairs:
            st = str(cr.status).lower()
            status_counts[st] += 1
            rnds = len(cr.negotiation_logs or [])
            base = Decimal(pm.base_rate)
            neg = Decimal(cr.negotiated_rate)
            prem = neg - base if neg > base else Decimal("0")
            rows.append(
                NegotiationRow(
                    contractor_rate_id=int(cr.id),
                    part_master_id=int(pm.id),
                    part_code=str(pm.part_code),
                    part_name=str(pm.part_name),
                    status=str(cr.status),
                    base_rate=base,
                    negotiated_rate=neg,
                    initial_rate=cr.initial_rate,
                    savings_amount=cr.savings_amount,
                    savings_percentage=cr.savings_percentage,
                    premium_above_base=prem,
                    rounds=rnds,
                    effective_from=cr.effective_from,
                    effective_to=cr.effective_to,
                    approved_at=cr.approved_at,
                )
            )
            bar_parts.append({"part_code": pm.part_code, "base_rate": float(base), "negotiated_rate": float(neg)})
            key = (cr.approved_at or cr.created_at).strftime("%Y-%m") if (cr.approved_at or cr.created_at) else "unknown"
            trend_map[key].append(neg)
            premium_rank.append((str(pm.part_code), prem))

        trend = [
            NegotiationTrendPoint(
                period=k,
                count=len(vals),
                avg_negotiated=(sum(vals) / len(vals)) if vals else None,
            )
            for k, vals in sorted(trend_map.items())
        ]
        premium_rank.sort(key=lambda x: x[1], reverse=True)
        pr_list = [{"part_code": a, "premium_above_base": float(b)} for a, b in premium_rank[:15]]

        return NegotiationAnalytics(
            rows=rows,
            status_counts=dict(status_counts),
            bar_chart_parts=bar_parts[:40],
            trend=trend[-24:],
            premium_ranking=pr_list,
        )

    def get_work_orders(self, contractor_id: int, flt: AnalyticsFilters) -> WorkOrderAnalytics:
        self._ensure(contractor_id)
        wo_base = select(WorkOrder).where(WorkOrder.contractor_id == int(contractor_id))
        wox = self._wo_date_filter(flt)
        if wox is not None:
            wo_base = wo_base.where(wox)
        wos = list(self._db.scalars(wo_base).all())

        totals_by_status: dict[str, int] = defaultdict(int)
        by_plant: dict[str, dict[str, Any]] = {}
        by_month_c: dict[str, int] = defaultdict(int)
        by_month_x: dict[str, int] = defaultdict(int)
        value_trend: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))

        for w in wos:
            st = str(w.status).lower()
            totals_by_status[st] += 1
            oid = int(w.org_unit_id)
            if oid not in by_plant:
                on = self._db.get(OrgUnit, oid)
                by_plant[oid] = {"plant_id": oid, "plant_name": on.name if on else str(oid), "count": 0, "value": Decimal("0")}
            by_plant[oid]["count"] += 1
            val = _d0(w.approved_value_total)
            sub = self._db.scalar(
                select(func.coalesce(func.sum(WorkOrderItem.taxable_value), 0)).where(WorkOrderItem.work_order_id == int(w.id))
            )
            if val == 0 and sub is not None:
                val = Decimal(str(sub))
            by_plant[oid]["value"] += val
            mkey = w.created_at.strftime("%Y-%m") if w.created_at else "unknown"
            by_month_c[mkey] += 1
            value_trend[mkey] += val
            if st == "closed" and w.updated_at:
                mx = w.updated_at.strftime("%Y-%m")
                by_month_x[mx] += 1

        months = sorted(set(by_month_c) | set(by_month_x) | set(value_trend.keys()))
        by_month_list = [
            WorkOrderMonthly(month=m, created=by_month_c.get(m, 0), completed=by_month_x.get(m, 0)) for m in months[-24:]
        ]
        vtrend = [{"month": m, "value": float(value_trend.get(m, Decimal("0")))} for m in months[-24:]]

        donut = [{"name": k, "value": v} for k, v in sorted(totals_by_status.items(), key=lambda kv: -kv[1])]

        inv_stmt = select(Invoice).where(Invoice.contractor_id == int(contractor_id))
        if flt.plant_id:
            inv_stmt = inv_stmt.where(Invoice.org_unit_id == int(flt.plant_id))
        invs = list(self._db.scalars(inv_stmt).all())
        total_invoiced = sum(_d0(i.total_amount) for i in invs if i.status in ("approved", "paid"))
        pending_inv = sum(_d0(i.total_amount) for i in invs if i.status in ("draft", "submitted", "pending_exception_approval", "blocked"))

        total_val = sum(Decimal(str(by_plant[p]["value"])) for p in by_plant)

        return WorkOrderAnalytics(
            totals_by_status=dict(totals_by_status),
            by_plant=list(by_plant.values()),
            by_month=by_month_list,
            total_wo_value=total_val,
            total_invoiced=total_invoiced,
            pending_invoice_amount=pending_inv,
            donut_status=donut,
            value_trend=vtrend,
        )

    def get_commercial(self, contractor_id: int, flt: AnalyticsFilters) -> CommercialInsights:
        neg = self.get_negotiations(contractor_id, flt)
        approved = [r for r in neg.rows if r.status == "approved"]
        if not approved:
            return CommercialInsights(
                avg_premium_above_base=None,
                avg_negotiation_reduction_pct=None,
                total_savings_generated=Decimal("0"),
                highest_premium_part_code=None,
                highest_premium_amount=None,
                highest_savings_part_code=None,
                highest_savings_amount=None,
                pct_negotiations_above_base=Decimal("0"),
                pct_negotiations_below_base=Decimal("0"),
                risk_level="low",
            )
        premiums = [r.premium_above_base for r in approved]
        avg_prem = sum(premiums) / len(premiums) if premiums else None
        sav_pcts = [_d0(r.savings_percentage) for r in approved if r.savings_percentage is not None]
        avg_red = (sum(sav_pcts) / len(sav_pcts)) if sav_pcts else None
        total_sav = sum(_d0(r.savings_amount) for r in approved)
        hp = max(approved, key=lambda r: r.premium_above_base)
        hs = max(approved, key=lambda r: _d0(r.savings_amount))
        above = sum(1 for r in approved if r.negotiated_rate > r.base_rate)
        below = sum(1 for r in approved if r.negotiated_rate < r.base_rate)
        n = len(approved)
        pct_above = Decimal(100 * above / n) if n else Decimal("0")
        pct_below = Decimal(100 * below / n) if n else Decimal("0")
        risk = "low"
        if avg_prem and avg_prem > Decimal("5") or pct_above > Decimal("60"):
            risk = "high"
        elif avg_prem and avg_prem > Decimal("2") or pct_above > Decimal("40"):
            risk = "medium"
        return CommercialInsights(
            avg_premium_above_base=avg_prem,
            avg_negotiation_reduction_pct=avg_red,
            total_savings_generated=total_sav,
            highest_premium_part_code=hp.part_code,
            highest_premium_amount=hp.premium_above_base,
            highest_savings_part_code=hs.part_code,
            highest_savings_amount=_d0(hs.savings_amount),
            pct_negotiations_above_base=pct_above,
            pct_negotiations_below_base=pct_below,
            risk_level=risk,
        )

    def get_pending(self, contractor_id: int) -> PendingActions:
        self._ensure(contractor_id)
        items: list[PendingItem] = []
        today = _today()
        warn_until = today + timedelta(days=30)

        for r in self._db.scalars(
            select(ContractorRate).where(
                ContractorRate.contractor_id == int(contractor_id),
                ContractorRate.status.in_(("draft", "pending_approval")),
            )
        ).all():
            sev = "warning" if r.status == "pending_approval" else "info"
            items.append(
                PendingItem(
                    kind="negotiation",
                    severity=sev,
                    title=f"Negotiation {r.status.replace('_', ' ')} — part #{r.part_master_id}",
                    detail="Submit, add rounds, or wait for approval.",
                    href_hint=f"/dashboard/negotiated-rates/{r.id}",
                    entity_id=int(r.id),
                )
            )

        for w in self._db.scalars(
            select(WorkOrder).where(
                WorkOrder.contractor_id == int(contractor_id),
                WorkOrder.status.in_(("draft", "pending_approval", "active")),
            )
        ).all():
            items.append(
                PendingItem(
                    kind="work_order",
                    severity="warning" if w.status != "active" else "info",
                    title=f"Work order {w.work_order_number} ({w.status})",
                    detail=w.title,
                    href_hint=f"/dashboard/work-orders/{w.id}",
                    entity_id=int(w.id),
                )
            )

        for inv in self._db.scalars(
            select(Invoice).where(
                Invoice.contractor_id == int(contractor_id),
                Invoice.status.in_(("draft", "submitted", "blocked", "pending_exception_approval")),
            )
        ).all():
            items.append(
                PendingItem(
                    kind="invoice",
                    severity="danger" if inv.status == "blocked" else "warning",
                    title=f"Invoice {inv.invoice_number} ({inv.status})",
                    entity_id=int(inv.id),
                    href_hint=f"/dashboard/invoices/{inv.id}",
                )
            )

        for r in self._db.scalars(
            select(ContractorRate).where(
                ContractorRate.contractor_id == int(contractor_id),
                ContractorRate.status == "approved",
                ContractorRate.effective_to.isnot(None),
                ContractorRate.effective_to >= today,
                ContractorRate.effective_to <= warn_until,
            )
        ).all():
            items.append(
                PendingItem(
                    kind="rate_expiry",
                    severity="warning",
                    title=f"Negotiated rate expiring soon (part #{r.part_master_id})",
                    detail=f"Effective to {r.effective_to}",
                    entity_id=int(r.id),
                    href_hint=f"/dashboard/negotiated-rates/{r.id}",
                )
            )

        for d in self._db.scalars(
            select(ContractorDocument).where(
                ContractorDocument.contractor_id == int(contractor_id),
                ContractorDocument.expiry_date.isnot(None),
                ContractorDocument.expiry_date >= today,
                ContractorDocument.expiry_date <= warn_until,
            )
        ).all():
            items.append(
                PendingItem(
                    kind="document",
                    severity="warning",
                    title=f"Document expiring: {d.document_name}",
                    detail=f"Expiry {d.expiry_date}",
                    entity_id=int(d.id),
                )
            )

        # Approval tasks for contractor rate approvals
        task_rows = list(
            self._db.execute(
                select(ApprovalTask, ApprovalRequest)
                .join(ApprovalRequest, ApprovalRequest.id == ApprovalTask.request_id)
                .join(ContractorRate, ContractorRate.id == ApprovalRequest.entity_id)
                .where(
                    ApprovalRequest.entity_type == "contractor_rate_approval",
                    ContractorRate.contractor_id == int(contractor_id),
                    ApprovalTask.status == "pending",
                )
            ).all()
        )
        for task, req in task_rows:
            items.append(
                PendingItem(
                    kind="approval_task",
                    severity="danger",
                    title=task.title or "Approval task pending",
                    detail=f"Request #{req.id}",
                    entity_id=int(task.id),
                    href_hint="/dashboard/tasks",
                )
            )

        return PendingActions(items=items)

    def get_timeline(self, contractor_id: int, categories: Sequence[str] | None) -> list[AnalyticsTimelineEvent]:
        self._ensure(contractor_id)
        cats = {str(c).lower() for c in categories} if categories else {"all"}
        use_all = "all" in cats
        out: list[AnalyticsTimelineEvent] = []

        def allow(cat: str) -> bool:
            return use_all or cat in cats

        # Work orders
        if allow("work_orders") or allow("financial"):
            for w in self._db.scalars(select(WorkOrder).where(WorkOrder.contractor_id == int(contractor_id))).all():
                if allow("work_orders"):
                    out.append(
                        AnalyticsTimelineEvent(
                            category="work_orders",
                            type="wo_created",
                            title=f"Work order {w.work_order_number}",
                            description=w.title,
                            timestamp=w.created_at,
                            metadata={"work_order_id": int(w.id), "status": w.status},
                        )
                    )
                if allow("financial") and w.status == "closed" and w.updated_at:
                    out.append(
                        AnalyticsTimelineEvent(
                            category="financial",
                            type="wo_closed",
                            title=f"Work order closed: {w.work_order_number}",
                            description=None,
                            timestamp=w.updated_at,
                            metadata={"work_order_id": int(w.id)},
                        )
                    )

        if allow("negotiations") or allow("approvals"):
            for cr in self._db.scalars(select(ContractorRate).where(ContractorRate.contractor_id == int(contractor_id))).all():
                pm = self._db.get(PartMaster, int(cr.part_master_id))
                label = f"{pm.part_code}" if pm else str(cr.part_master_id)
                if allow("negotiations"):
                    out.append(
                        AnalyticsTimelineEvent(
                            category="negotiations",
                            type="rate_row",
                            title=f"Negotiation record — {label}",
                            description=f"Status {cr.status}",
                            timestamp=cr.created_at,
                            metadata={"contractor_rate_id": int(cr.id), "status": cr.status},
                        )
                    )
                for log in cr.negotiation_logs or []:
                    if allow("negotiations"):
                        out.append(
                            AnalyticsTimelineEvent(
                                category="negotiations",
                                type="negotiation_round",
                                title=f"Round {log.round_number} — {label}",
                                description=log.remarks,
                                timestamp=log.created_at,
                                metadata={"round": log.round_number, "contractor_rate_id": int(cr.id)},
                            )
                        )
                if allow("approvals") and cr.approved_at:
                    out.append(
                        AnalyticsTimelineEvent(
                            category="approvals",
                            type="rate_approved",
                            title=f"Negotiation approved — {label}",
                            description=None,
                            timestamp=cr.approved_at,
                            metadata={"contractor_rate_id": int(cr.id)},
                        )
                    )

        if allow("financial"):
            for inv in self._db.scalars(select(Invoice).where(Invoice.contractor_id == int(contractor_id))).all():
                out.append(
                    AnalyticsTimelineEvent(
                        category="financial",
                        type="invoice",
                        title=f"Invoice {inv.invoice_number} ({inv.status})",
                        description=None,
                        timestamp=inv.created_at,
                        metadata={"invoice_id": int(inv.id), "total": str(inv.total_amount)},
                    )
                )

        if allow("compliance"):
            for d in self._db.scalars(select(ContractorDocument).where(ContractorDocument.contractor_id == int(contractor_id))).all():
                out.append(
                    AnalyticsTimelineEvent(
                        category="compliance",
                        type="document",
                        title=f"Document: {d.document_name}",
                        description=f"Verification {d.verification_status}",
                        timestamp=d.created_at,
                        metadata={"document_id": int(d.id)},
                    )
                )

        out.sort(key=lambda e: e.timestamp, reverse=True)
        return out[:500]

    def build_xlsx(self, contractor_id: int, flt: AnalyticsFilters) -> bytes:
        try:
            from openpyxl import Workbook
        except ImportError as e:  # pragma: no cover
            raise RuntimeError("openpyxl is required for Excel export. pip install openpyxl") from e

        self._ensure(contractor_id)
        wb = Workbook()
        # Negotiations
        ws1 = wb.active
        ws1.title = "Negotiations"
        neg = self.get_negotiations(contractor_id, flt)
        headers = [
            "contractor_rate_id",
            "part_code",
            "part_name",
            "status",
            "base_rate",
            "negotiated_rate",
            "initial_rate",
            "savings_amount",
            "savings_pct",
            "premium_above_base",
            "rounds",
            "effective_from",
            "effective_to",
        ]
        ws1.append(headers)
        for r in neg.rows:
            ws1.append(
                [
                    r.contractor_rate_id,
                    r.part_code,
                    r.part_name,
                    r.status,
                    float(r.base_rate),
                    float(r.negotiated_rate),
                    float(r.initial_rate) if r.initial_rate is not None else None,
                    float(r.savings_amount) if r.savings_amount is not None else None,
                    float(r.savings_percentage) if r.savings_percentage is not None else None,
                    float(r.premium_above_base),
                    r.rounds,
                    r.effective_from.isoformat(),
                    r.effective_to.isoformat() if r.effective_to else None,
                ]
            )

        ws2 = wb.create_sheet("Work orders")
        wo = self.get_work_orders(contractor_id, flt)
        ws2.append(["status", "count"])
        for k, v in sorted(wo.totals_by_status.items()):
            ws2.append([k, v])
        ws2.append([])
        ws2.append(["plant", "count", "value"])
        for p in wo.by_plant:
            ws2.append([p.get("plant_name"), p.get("count"), float(p.get("value", 0))])

        ws3 = wb.create_sheet("Invoices")
        ws3.append(["id", "invoice_number", "status", "total_amount", "invoice_date"])
        for inv in self._db.scalars(select(Invoice).where(Invoice.contractor_id == int(contractor_id))).all():
            ws3.append(
                [
                    int(inv.id),
                    inv.invoice_number,
                    inv.status,
                    float(inv.total_amount),
                    inv.invoice_date.isoformat(),
                ]
            )

        ws5 = wb.create_sheet("Timeline")
        tl = self.get_timeline(contractor_id, ("all",))
        ws5.append(["timestamp", "category", "type", "title"])
        for e in tl[:300]:
            ws5.append([e.timestamp.isoformat(), e.category, e.type, e.title])

        ws6 = wb.create_sheet("Commercial")
        com = self.get_commercial(contractor_id, flt)
        ws6.append(["metric", "value"])
        ws6.append(["avg_premium_above_base", float(com.avg_premium_above_base) if com.avg_premium_above_base else None])
        ws6.append(["avg_negotiation_reduction_pct", float(com.avg_negotiation_reduction_pct) if com.avg_negotiation_reduction_pct else None])
        ws6.append(["total_savings_generated", float(com.total_savings_generated)])
        ws6.append(["risk_level", com.risk_level])

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
